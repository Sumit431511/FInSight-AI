from pathlib import Path
import hashlib
import sqlite3
import os
import pandas as pd
from threading import Lock
from dotenv import load_dotenv

from app.config import get_data_dir
from app.config import get_groq_model


class SimpleDocument:
    def __init__(self, page_content, metadata=None):
        self.page_content = page_content
        self.metadata = metadata or {}

try:
    from langchain_core.documents import Document
    from langchain_core.prompts import ChatPromptTemplate
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    _IMPORT_ERROR = None
except ImportError as exc:
    Document = None
    ChatPromptTemplate = None
    RecursiveCharacterTextSplitter = None
    _IMPORT_ERROR = exc

Chroma = None
HuggingFaceEmbeddings = None
ChatGroq = None
create_stuff_documents_chain = None
create_retrieval_chain = None
CohereRerank = None
ContextualCompressionRetriever = None

if Document is None:
    Document = SimpleDocument

load_dotenv()

LANGCHAIN_API_KEY = os.getenv("LANGCHAIN_API_KEY")
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
COHERE_API_KEY = os.getenv("COHERE_API_KEY")

os.environ["LANGCHAIN_TRACING_V2"] = "true"
os.environ["LANGCHAIN_ENDPOINT"] = "https://api.smith.langchain.com"
os.environ["LANGCHAIN_PROJECT"] = "RBAC-RAG"

embeddings = None
vectorstore = None
question_answering_chain = None
fallback_documents = []
indexer_lock = Lock()


class _FallbackRAGChain:
    def __init__(self, user_role: str):
        self.user_role = user_role.lower()

    def invoke(self, payload):
        question = payload.get("input", "")
        docs = [
            doc for doc in fallback_documents
            if _doc_is_visible_for_role(doc, self.user_role)
        ]

        if not docs:
            return {
                "answer": "No indexed documents were available for your role yet.",
                "context": [],
            }

        context_text = "\n\n".join(
            f"Source: {doc.metadata.get('source', 'unknown')}\n{doc.page_content[:800]}"
            for doc in docs
        )

        answer = (
            f"I found {len(docs)} relevant document(s) for your role.\n\n"
            f"{context_text}\n\n"
            f"Question: {question}"
        )

        return {"answer": answer, "context": docs}


def _doc_is_visible_for_role(doc, user_role: str) -> bool:
    role = (doc.metadata.get("role") or "").lower()
    if user_role == "c-level":
        return True
    if user_role == "general":
        return role in {"", "general"}
    return role in {"", user_role, "general"}


def _ensure_rag_dependencies():
    if _IMPORT_ERROR is not None:
        return True

    global Chroma, HuggingFaceEmbeddings, ChatGroq, create_stuff_documents_chain
    global create_retrieval_chain, CohereRerank, ContextualCompressionRetriever
    global embeddings, vectorstore, question_answering_chain

    try:
        if Chroma is None:
            from langchain_chroma import Chroma as _Chroma
            from langchain_huggingface import HuggingFaceEmbeddings as _HuggingFaceEmbeddings
            from langchain_groq import ChatGroq as _ChatGroq
            from langchain_classic.chains.combine_documents import create_stuff_documents_chain as _csdc
            from langchain_classic.chains.retrieval import create_retrieval_chain as _crc
            from langchain_cohere import CohereRerank as _CohereRerank
            from langchain_classic.retrievers import ContextualCompressionRetriever as _ccr

            Chroma = _Chroma
            HuggingFaceEmbeddings = _HuggingFaceEmbeddings
            ChatGroq = _ChatGroq
            create_stuff_documents_chain = _csdc
            create_retrieval_chain = _crc
            CohereRerank = _CohereRerank
            ContextualCompressionRetriever = _ccr

        if embeddings is None:
            embeddings = HuggingFaceEmbeddings(
                model_name="BAAI/bge-small-en-v1.5",
                model_kwargs={"device": "cpu"},
            )

        if vectorstore is None:
            vectorstore = Chroma(
                collection_name="my_collection",
                persist_directory=str(get_data_dir("chroma_db")),
                embedding_function=embeddings,
            )

        if question_answering_chain is None:
            system_prompt = (
                "You are an assistant for summarizing and answering queries from internal company documents.\n"
                "Always use the retrieved context to answer the query, even if partial.\n"
                "Do not guess. If data is not found, explain what you searched for.\n"
                "When responding:\n"
                "- Add **Source** from document metadata if possible.\n"
                "- Use headers\n"
                "- Use bullet points\n"
                "- For CSV-style data, format in table with two columns\n"
                "\n{context}"
            )
            chat_prompt = ChatPromptTemplate.from_messages([
                ("system", system_prompt),
                ("human", "{input}"),
            ])
            model = ChatGroq(model=get_groq_model(), temperature=0.2)
            question_answering_chain = create_stuff_documents_chain(model, chat_prompt)
    except Exception as exc:
        print(f"RAG runtime initialization failed; using in-memory fallback: {exc}")
        return False

    return True


def embed_documents_to_vectorstore(docs, document_id: int | None = None):
    if not docs:
        return True

    normalized_docs = []
    if isinstance(docs, list):
        normalized_docs = docs
    else:
        normalized_docs = [docs]

    if _IMPORT_ERROR is not None:
        fallback_documents.extend(normalized_docs)
        print("Using in-memory fallback storage because optional RAG dependencies are unavailable.")
        return True

    if not _ensure_rag_dependencies():
        fallback_documents.extend(normalized_docs)
        return True

    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    splits = text_splitter.split_documents(normalized_docs)
    ids = []
    for index, split in enumerate(splits):
        stable_document_id = str(split.metadata.get("document_id", document_id or "unknown"))
        content_hash = hashlib.sha256(split.page_content.encode("utf-8")).hexdigest()[:16]
        ids.append(f"{stable_document_id}:{index}:{content_hash}")
    vectorstore.add_documents(splits, ids=ids)

    print("Documents embedded and saved to vectorstore.")
    print("Total documents:", len(vectorstore.get()["documents"]))
    return True

def load_file(
    filepath,
    role,
    document_id: int | None = None,
    source_name: str | None = None,
):
    ext = Path(filepath).suffix.lower()
    source_name = source_name or Path(filepath).name
    try:
        if ext == ".csv":
            df1 = pd.read_csv(filepath)
            documents = []
            for row in df1.to_dict(orient="records"):
                content = "\n".join(f"{k}: {v}" for k, v in row.items())
                documents.append(
                    Document(
                        page_content=content,
                        metadata={"role": role.lower(), "source": source_name}
                    )
                )
            for document in documents:
                document.metadata["document_id"] = str(document_id) if document_id is not None else ""
            return documents

        elif ext in {".md", ".txt", ".text"}:
            with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return [
                Document(
                    page_content=content,
                    metadata={
                        "role": role.lower(),
                        "source": source_name,
                        "document_id": str(document_id) if document_id is not None else "",
                    }
                )
            ]

        elif ext == ".pdf":
            try:
                from langchain_community.document_loaders import PyPDFLoader
                loader = PyPDFLoader(str(filepath))
                pages = loader.load()
                for page in pages:
                    page.metadata["role"] = role.lower()
                    page.metadata["source"] = source_name
                    page.metadata["document_id"] = str(document_id) if document_id is not None else ""
                return pages
            except Exception:
                import pypdf
                reader = pypdf.PdfReader(str(filepath))
                content = "\n".join([page.extract_text() or "" for page in reader.pages])
                return [
                    Document(
                        page_content=content,
                        metadata={
                            "role": role.lower(),
                            "source": source_name,
                            "document_id": str(document_id) if document_id is not None else "",
                        }
                    )
                ]

        elif ext == ".docx":
            import docx
            doc = docx.Document(filepath)
            content = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            return [
                Document(
                    page_content=content,
                    metadata={
                        "role": role.lower(),
                        "source": source_name,
                        "document_id": str(document_id) if document_id is not None else "",
                    }
                )
            ]

        else:
            return None

    except Exception as e:
        print(f"Failed to process {filepath}: {e}")
        return None


from app.db import get_sqlite_conn


def run_indexer():
    """Index pending documents and persist per-document success or failure state."""
    if not indexer_lock.acquire(blocking=False):
        return {"indexed": [], "failed": [], "skipped": True}

    indexed, failed = [], []
    try:
        with get_sqlite_conn() as conn:
            pending = conn.execute(
                "SELECT id, filepath, role, filename FROM documents WHERE status = 'pending'"
            ).fetchall()

        for row in pending:
            document_id, path, role, filename = row["id"], row["filepath"], row["role"], row["filename"]
            try:
                with get_sqlite_conn() as conn:
                    conn.execute(
                        "UPDATE documents SET status = 'indexing', error_message = NULL WHERE id = ?",
                        (document_id,),
                    )
                    conn.commit()

                docs = load_file(
                    path,
                    role,
                    document_id=document_id,
                    source_name=filename,
                )
                if not docs or not embed_documents_to_vectorstore(docs, document_id=document_id):
                    raise ValueError("The document could not be converted into searchable content.")

                with get_sqlite_conn() as conn:
                    conn.execute(
                        "UPDATE documents SET embedded = 1, status = 'indexed', error_message = NULL WHERE id = ?",
                        (document_id,),
                    )
                    conn.commit()
                indexed.append(document_id)
            except Exception as exc:
                with get_sqlite_conn() as conn:
                    conn.execute(
                        "UPDATE documents SET embedded = 0, status = 'failed', error_message = ? WHERE id = ?",
                        (str(exc)[:500], document_id),
                    )
                    conn.commit()
                failed.append(document_id)
        return {"indexed": indexed, "failed": failed, "skipped": False}
    finally:
        indexer_lock.release()


def delete_document_vectors(document_id: int) -> None:
    """Delete all vectors associated with one document without affecting other roles."""
    if _IMPORT_ERROR is not None:
        fallback_documents[:] = [
            document for document in fallback_documents
            if document.metadata.get("document_id") != str(document_id)
        ]
        return

    _ensure_rag_dependencies()
    vectorstore.delete(where={"document_id": str(document_id)})

def wrap_with_reranker(retriever, cohere_api_key, top_n=4):
    _ensure_rag_dependencies()
    reranker = CohereRerank(cohere_api_key=cohere_api_key, top_n=top_n)
    return ContextualCompressionRetriever(
        base_compressor=reranker,
        base_retriever=retriever,
    )


def get_rag_chain(user_role: str, cohere_api_key: str = None):
    if _IMPORT_ERROR is not None:
        return _FallbackRAGChain(user_role)

    if not _ensure_rag_dependencies():
        return _FallbackRAGChain(user_role)

    user_role = user_role.lower()

    if user_role == "c-level":
        retriever = vectorstore.as_retriever(search_kwargs={"k": 4})

    elif user_role == "general":
        retriever = vectorstore.as_retriever(search_kwargs={
            "k": 4,
            "filter": {"role": "general"},
        })

    else:
        retriever = vectorstore.as_retriever(search_kwargs={
            "k": 4,
            "filter": {"role": {"$in": [user_role, "general"]}},
        })

    if cohere_api_key:
        print("Using cohere reranker")
        retriever = wrap_with_reranker(retriever, cohere_api_key)

    return create_retrieval_chain(retriever, question_answering_chain)
    """
    from langchain_core.runnables import RunnableLambda, RunnableMap

    extract_input = RunnableLambda(lambda x: x["input"])

    return RunnableMap({
        "context": extract_input | retriever,
        "answer": extract_input | retriever | question_answering_chain
    })"""


"""
# ========== MAIN EXECUTION ==========
if __name__ == "__main__":
    run_indexer() 
"""
    
"""
    user_role = "hr" 
    rag_chain = get_rag_chain(user_role)

    
    query = "give me Campaign Highlights from marketing summary."
    response = rag_chain.invoke({"input": query})

    print((response["answer"]))
    for doc in response.get("context", []):
        print(f"Source: {doc.metadata['source']}, Role: {doc.metadata.get('role')}")

"""

